"""Text extraction for common, non-executable document formats."""

from __future__ import annotations

import io
import zipfile

MAX_EXTRACTED_CHARS = 100_000


def extract_text(data: bytes, suffix: str) -> str:
    """Return bounded searchable text without executing uploaded content."""
    extractors = {
        ".txt": _plain_text,
        ".md": _plain_text,
        ".csv": _plain_text,
        ".pdf": _pdf,
        ".doc": _ole_text,
        ".docx": _docx,
        ".ppt": _ole_text,
        ".pptx": _pptx,
        ".xls": _xls,
        ".xlsx": _xlsx,
        ".odt": _odf,
        ".odp": _odf,
        ".ods": _odf,
    }
    extractor = extractors.get(suffix.lower())
    if extractor is None:
        return ""
    try:
        return extractor(data)[:MAX_EXTRACTED_CHARS]
    except (OSError, ValueError, KeyError, UnicodeError, zipfile.BadZipFile):
        return ""


def _plain_text(data: bytes) -> str:
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    return "\n".join(
        page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages
    )



def _docx(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    return (
        "\n".join(paragraph.text for paragraph in document.paragraphs)
        + "\n"
        + "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
    )


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    return "\n".join(
        shape.text  # pyright: ignore[reportAttributeAccessIssue]
        for slide in Presentation(io.BytesIO(data)).slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return "\n".join(
            " ".join(str(value) for value in row if value is not None)
            for sheet in workbook.worksheets
            for row in sheet.iter_rows(values_only=True)
        )
    finally:
        workbook.close()


def _xls(data: bytes) -> str:
    import xlrd

    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        return "\n".join(
            " ".join(
                str(sheet.cell_value(row, column))
                for column in range(sheet.ncols)
                if sheet.cell_value(row, column) != ""
            )
            for sheet in (
                workbook.sheet_by_index(index) for index in range(workbook.nsheets)
            )
            for row in range(sheet.nrows)
        )
    finally:
        workbook.release_resources()


def _ole_text(data: bytes) -> str:
    """Extract printable strings from legacy Word/PowerPoint OLE containers."""
    import re
    import olefile

    with olefile.OleFileIO(io.BytesIO(data)) as document:
        blobs = []
        for stream in document.listdir(streams=True, storages=False):
            with document.openstream(stream) as source:
                blobs.append(source.read())
    raw = b"\n".join(blobs)
    decoded = (
        raw.decode("utf-16-le", errors="ignore")
        + "\n"
        + raw.decode("latin-1", errors="ignore")
    )
    return "\n".join(re.findall(r"[\wÀ-ÿ][\wÀ-ÿ\s.,;:!?()'\"/-]{3,}", decoded))


def _odf(data: bytes) -> str:
    from odf import teletype
    from odf.opendocument import load

    document = load(io.BytesIO(data))
    return teletype.extractText(document)
